from __future__ import annotations

import os
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .artifacts import ArtifactDefinition
from .config import CONTROL_FILE

BASE_COLUMNS = [
    "FILE_NAME", "FILE_PATH", "SOURCE_KIND", "FILE_SIZE", "FILE_HASH", "LAST_MODIFIED",
    "TXT_STATUS", "TXT_LAST_RUN", "TXT_ERROR", "TXT_HASH",
    "EMBEDDING_STATUS", "EMBEDDING_LAST_RUN", "EMBEDDING_ERROR", "EMBEDDING_TXT_HASH",
    "XML_STATUS", "XML_LAST_RUN", "XML_ERROR", "XML_SDMX_VERSION",
    "CONSISTENCY_STATUS", "CONSISTENCY_LAST_RUN", "CONSISTENCY_ERROR", "CONSISTENCY_PASSES",
]
STATUSES = {"PENDING", "RUNNING", "SUCCESS", "FAILED", "SKIPPED", "MODIFIED"}
AGGREGATE_COLUMNS = [
    "ARTIFACT_KEY", "ARTIFACT_NAME", "STATUS", "LAST_RUN", "ERROR",
    "AGGREGATE_JSON_PATH", "AGGREGATE_JSON_HASH", "JSON_MTIME_NS",
    "CSV_MTIME_NS", "CONTRIBUTION_COUNT", "PRINCIPLES_HASH",
    "FINAL_STATUS", "FINAL_LAST_RUN", "FINAL_ERROR", "FINAL_INPUT_HASH",
    "FINAL_JSON_PATH", "FINAL_JSON_HASH", "FINAL_JSON_MTIME_NS", "FINAL_CSV_PATH",
    "FINAL_REVIEW_ATTEMPTS", "FINAL_VALIDATOR_CORRECTION",
]
XML_EXPORT_COLUMNS = [
    "SDMX_VERSION", "STATUS", "LAST_RUN", "ERROR", "INPUT_HASH",
    "XML_PATH", "XML_HASH", "XML_MTIME_NS",
]


def now() -> str:
    return datetime.now(timezone.utc).isoformat()


class ControlWorkbook:
    def __init__(self, artifacts: list[ArtifactDefinition]):
        self.artifacts = artifacts
        if CONTROL_FILE.exists():
            self.workbook = load_workbook(CONTROL_FILE)
            self.sheet = self.workbook["Pipeline Control"]
        else:
            self.workbook = Workbook()
            self.sheet = self.workbook.active
            self.sheet.title = "Pipeline Control"
        self._ensure_columns()
        self._ensure_aggregate_sheet()
        self._ensure_xml_export_sheet()

    @property
    def columns(self) -> list[str]:
        return [self.sheet.cell(1, column).value for column in range(1, self.sheet.max_column + 1)]

    def _ensure_columns(self) -> None:
        while self.sheet.max_column > 1 and self.sheet.cell(1, 1).value is None:
            self.sheet.delete_cols(1)
        existing = [self.sheet.cell(1, column).value for column in range(1, self.sheet.max_column + 1)] if self.sheet.max_row else []
        required = list(BASE_COLUMNS)
        for artifact in self.artifacts:
            required.extend([
                f"{artifact.key}_STATUS", f"{artifact.key}_LAST_RUN", f"{artifact.key}_ERROR",
                f"{artifact.key}_PRINCIPLES_HASH", f"{artifact.key}_INPUT_HASH",
                f"{artifact.key}_JSON_PATH", f"{artifact.key}_JSON_MTIME_NS",
                f"{artifact.key}_CSV_STATUS", f"{artifact.key}_CSV_LAST_RUN",
                f"{artifact.key}_CSV_ERROR", f"{artifact.key}_CSV_MTIME_NS",
                f"{artifact.key}_AGGREGATED", f"{artifact.key}_AGGREGATED_JSON_HASH",
                f"{artifact.key}_AGGREGATED_LAST_RUN",
                f"{artifact.key}_SEMANTIC_ATTEMPTS", f"{artifact.key}_VALIDATOR_CORRECTION",
            ])
        for name in required:
            if name not in existing:
                existing.append(name)
                self.sheet.cell(1, len(existing), name)
        consistency_column = existing.index("CONSISTENCY_STATUS") + 1
        source_kind_column = existing.index("SOURCE_KIND") + 1
        file_path_column = existing.index("FILE_PATH") + 1
        xml_status_column = existing.index("XML_STATUS") + 1
        for row in range(2, self.sheet.max_row + 1):
            if self.sheet.cell(row, consistency_column).value is None:
                self.sheet.cell(row, consistency_column, "PENDING")
            if self.sheet.cell(row, source_kind_column).value is None:
                is_xml = str(self.sheet.cell(row, file_path_column).value or "").replace("\\", "/").startswith("Inputs XML/")
                self.sheet.cell(row, source_kind_column, "XML" if is_xml else "REPORT")
            if self.sheet.cell(row, xml_status_column).value is None and self.sheet.cell(row, source_kind_column).value == "REPORT":
                self.sheet.cell(row, xml_status_column, "SKIPPED")
        for cell in self.sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        self.sheet.freeze_panes = "A2"
        self.sheet.auto_filter.ref = self.sheet.dimensions
        for index, name in enumerate(existing, 1):
            self.sheet.column_dimensions[get_column_letter(index)].width = min(40, max(12, len(str(name)) + 2))
        self.save()

    def _ensure_aggregate_sheet(self) -> None:
        if "Aggregate Control" in self.workbook.sheetnames:
            sheet = self.workbook["Aggregate Control"]
        else:
            sheet = self.workbook.create_sheet("Aggregate Control")
        for index, name in enumerate(AGGREGATE_COLUMNS, 1):
            sheet.cell(1, index, name)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="38761D")
        sheet.freeze_panes = "A2"
        for artifact in self.artifacts:
            row = self.aggregate_row(artifact.key, create=False)
            if row is None:
                row = sheet.max_row + 1
                sheet.cell(row, 1, artifact.key)
                sheet.cell(row, 2, artifact.name)
                sheet.cell(row, 3, "PENDING")
            else:
                sheet.cell(row, 2, artifact.name)
        for index, name in enumerate(AGGREGATE_COLUMNS, 1):
            sheet.column_dimensions[get_column_letter(index)].width = min(45, max(14, len(name) + 2))
        self.aggregate_sheet = sheet
        self.save()

    def _ensure_xml_export_sheet(self) -> None:
        if "XML Export Control" in self.workbook.sheetnames:
            sheet = self.workbook["XML Export Control"]
        else:
            sheet = self.workbook.create_sheet("XML Export Control")
        for index, name in enumerate(XML_EXPORT_COLUMNS, 1):
            sheet.cell(1, index, name)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="674EA7")
        sheet.freeze_panes = "A2"
        versions = {sheet.cell(row, 1).value: row for row in range(2, sheet.max_row + 1)}
        for version in ("2.1", "3.0"):
            if version not in versions:
                row = sheet.max_row + 1
                sheet.cell(row, 1, version)
                sheet.cell(row, 2, "PENDING")
        for index, name in enumerate(XML_EXPORT_COLUMNS, 1):
            sheet.column_dimensions[get_column_letter(index)].width = min(45, max(14, len(name) + 2))
        self.xml_export_sheet = sheet
        self.save()

    def save(self) -> None:
        temporary = Path(str(CONTROL_FILE) + f".{os.getpid()}.tmp.xlsx")
        self.workbook.save(temporary)
        temporary.replace(CONTROL_FILE)

    def _row_map(self, row: int) -> dict[str, Any]:
        return {name: self.sheet.cell(row, index + 1).value for index, name in enumerate(self.columns)}

    def find_row(self, relative_path: str) -> int | None:
        path_column = self.columns.index("FILE_PATH") + 1
        for row in range(2, self.sheet.max_row + 1):
            if self.sheet.cell(row, path_column).value == relative_path:
                return row
        return None

    def sync_file(self, source: Path, relative_path: str, file_hash: str) -> tuple[int, bool]:
        row = self.find_row(relative_path)
        is_new = row is None
        if row is None:
            row = self.sheet.max_row + 1
            for name in ("TXT_STATUS", "EMBEDDING_STATUS", "CONSISTENCY_STATUS"):
                self.set(row, name, "PENDING", save=False)
            for artifact in self.artifacts:
                self.set(row, f"{artifact.key}_STATUS", "PENDING", save=False)
                self.set(row, f"{artifact.key}_CSV_STATUS", "PENDING", save=False)
                self.set(row, f"{artifact.key}_AGGREGATED", "NO", save=False)
        previous_hash = self.get(row, "FILE_HASH")
        changed = previous_hash not in (None, file_hash)
        self.set(row, "FILE_NAME", source.name, save=False)
        self.set(row, "FILE_PATH", relative_path, save=False)
        self.set(row, "SOURCE_KIND", "REPORT", save=False)
        self.set(row, "XML_STATUS", "SKIPPED", save=False)
        self.set(row, "FILE_SIZE", source.stat().st_size, save=False)
        self.set(row, "FILE_HASH", file_hash, save=False)
        self.set(row, "LAST_MODIFIED", datetime.fromtimestamp(source.stat().st_mtime, timezone.utc).isoformat(), save=False)
        if changed:
            for prefix in ["TXT", "EMBEDDING", "CONSISTENCY", *[artifact.key for artifact in self.artifacts]]:
                self.set(row, f"{prefix}_STATUS", "PENDING", save=False)
                self.set(row, f"{prefix}_ERROR", "", save=False)
            for artifact in self.artifacts:
                self.set(row, f"{artifact.key}_CSV_STATUS", "PENDING", save=False)
                self.set(row, f"{artifact.key}_AGGREGATED", "NO", save=False)
                self.set(row, f"{artifact.key}_AGGREGATED_JSON_HASH", "", save=False)
        for artifact in self.artifacts:
            old_principles = self.get(row, f"{artifact.key}_PRINCIPLES_HASH")
            if old_principles not in (None, artifact.sha256):
                self.set(row, f"{artifact.key}_STATUS", "PENDING", save=False)
                self.set(row, f"{artifact.key}_ERROR", "", save=False)
                self.set(row, f"{artifact.key}_AGGREGATED", "NO", save=False)
                self.set(row, f"{artifact.key}_AGGREGATED_JSON_HASH", "", save=False)
        self.save()
        return row, is_new or changed

    def sync_xml_file(self, source: Path, relative_path: str, file_hash: str, sdmx_version: str) -> tuple[int, bool]:
        row = self.find_row(relative_path)
        is_new = row is None
        if row is None:
            row = self.sheet.max_row + 1
            for artifact in self.artifacts:
                self.set(row, f"{artifact.key}_STATUS", "PENDING", save=False)
                self.set(row, f"{artifact.key}_CSV_STATUS", "PENDING", save=False)
                self.set(row, f"{artifact.key}_AGGREGATED", "NO", save=False)
        previous_hash = self.get(row, "FILE_HASH")
        changed = previous_hash not in (None, file_hash)
        self.set(row, "FILE_NAME", source.name, save=False)
        self.set(row, "FILE_PATH", relative_path, save=False)
        self.set(row, "SOURCE_KIND", "XML", save=False)
        self.set(row, "FILE_SIZE", source.stat().st_size, save=False)
        self.set(row, "FILE_HASH", file_hash, save=False)
        self.set(row, "LAST_MODIFIED", datetime.fromtimestamp(source.stat().st_mtime, timezone.utc).isoformat(), save=False)
        self.set(row, "TXT_STATUS", "SKIPPED", save=False)
        self.set(row, "EMBEDDING_STATUS", "SKIPPED", save=False)
        self.set(row, "XML_SDMX_VERSION", sdmx_version, save=False)
        if is_new or changed:
            self.set(row, "XML_STATUS", "PENDING", save=False)
            self.set(row, "CONSISTENCY_STATUS", "PENDING", save=False)
            for artifact in self.artifacts:
                self.set(row, f"{artifact.key}_STATUS", "PENDING", save=False)
                self.set(row, f"{artifact.key}_AGGREGATED", "NO", save=False)
                self.set(row, f"{artifact.key}_AGGREGATED_JSON_HASH", "", save=False)
        self.save()
        return row, is_new or changed

    def get(self, row: int, column: str) -> Any:
        return self.sheet.cell(row, self.columns.index(column) + 1).value

    def set(self, row: int, column: str, value: Any, save: bool = True) -> None:
        if column.endswith("_STATUS") and value not in STATUSES:
            raise ValueError(f"Invalid status: {value}")
        self.sheet.cell(row, self.columns.index(column) + 1, value)
        if save:
            self.save()

    def stage(self, row: int, prefix: str, status: str, error: str = "") -> None:
        self.set(row, f"{prefix}_STATUS", status, save=False)
        if f"{prefix}_LAST_RUN" in self.columns:
            self.set(row, f"{prefix}_LAST_RUN", now(), save=False)
        if f"{prefix}_ERROR" in self.columns:
            self.set(row, f"{prefix}_ERROR", error[:30000], save=False)
        self.save()

    def data_rows(self) -> list[int]:
        return list(range(2, self.sheet.max_row + 1))

    def aggregate_row(self, artifact_key: str, create: bool = True) -> int | None:
        sheet = self.workbook["Aggregate Control"] if "Aggregate Control" in self.workbook.sheetnames else None
        if sheet is not None:
            for row in range(2, sheet.max_row + 1):
                if sheet.cell(row, 1).value == artifact_key:
                    return row
        if not create:
            return None
        if sheet is None:
            sheet = self.workbook.create_sheet("Aggregate Control")
            for index, name in enumerate(AGGREGATE_COLUMNS, 1):
                sheet.cell(1, index, name)
        row = sheet.max_row + 1
        sheet.cell(row, 1, artifact_key)
        sheet.cell(row, 3, "PENDING")
        self.aggregate_sheet = sheet
        return row

    def aggregate_get(self, artifact_key: str, column: str) -> Any:
        row = self.aggregate_row(artifact_key)
        return self.aggregate_sheet.cell(row, AGGREGATE_COLUMNS.index(column) + 1).value

    def aggregate_set(self, artifact_key: str, column: str, value: Any, save: bool = True) -> None:
        row = self.aggregate_row(artifact_key)
        self.aggregate_sheet.cell(row, AGGREGATE_COLUMNS.index(column) + 1, value)
        if save:
            self.save()

    def final_stage(self, artifact_key: str, status: str, error: str = "") -> None:
        if status not in STATUSES:
            raise ValueError(f"Invalid final status: {status}")
        self.aggregate_set(artifact_key, "FINAL_STATUS", status, save=False)
        self.aggregate_set(artifact_key, "FINAL_LAST_RUN", now(), save=False)
        self.aggregate_set(artifact_key, "FINAL_ERROR", error[:30000], save=False)
        self.save()

    def aggregate_stage(self, artifact_key: str, status: str, error: str = "") -> None:
        if status not in STATUSES:
            raise ValueError(f"Invalid aggregate status: {status}")
        self.aggregate_set(artifact_key, "STATUS", status, save=False)
        self.aggregate_set(artifact_key, "LAST_RUN", now(), save=False)
        self.aggregate_set(artifact_key, "ERROR", error[:30000], save=False)
        self.save()

    def xml_export_row(self, version: str) -> int:
        for row in range(2, self.xml_export_sheet.max_row + 1):
            if str(self.xml_export_sheet.cell(row, 1).value) == version:
                return row
        row = self.xml_export_sheet.max_row + 1
        self.xml_export_sheet.cell(row, 1, version)
        self.xml_export_sheet.cell(row, 2, "PENDING")
        return row

    def xml_export_get(self, version: str, column: str) -> Any:
        return self.xml_export_sheet.cell(self.xml_export_row(version), XML_EXPORT_COLUMNS.index(column) + 1).value

    def xml_export_set(self, version: str, column: str, value: Any, save: bool = True) -> None:
        self.xml_export_sheet.cell(self.xml_export_row(version), XML_EXPORT_COLUMNS.index(column) + 1, value)
        if save:
            self.save()

    def xml_export_stage(self, version: str, status: str, error: str = "") -> None:
        if status not in STATUSES:
            raise ValueError(f"Invalid XML export status: {status}")
        self.xml_export_set(version, "STATUS", status, save=False)
        self.xml_export_set(version, "LAST_RUN", now(), save=False)
        self.xml_export_set(version, "ERROR", error[:30000], save=False)
        self.save()

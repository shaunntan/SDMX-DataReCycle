from __future__ import annotations

import csv
import html
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

import fitz
import pandas as pd
from bs4 import BeautifulSoup
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation


def _table(rows: list[list[object]]) -> str:
    cleaned = [["" if value is None else str(value).replace("\n", " ").strip() for value in row] for row in rows]
    width = max((len(row) for row in cleaned), default=0)
    if not width:
        return ""
    padded = [row + [""] * (width - len(row)) for row in cleaned]
    return "\n".join("| " + " | ".join(row) + " |" for row in padded)


def _office_legacy(path: Path, target_extension: str) -> Path:
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if not executable:
        raise RuntimeError(f"{path.suffix.upper()} conversion requires free LibreOffice (soffice) on PATH")
    temp_dir = Path(tempfile.mkdtemp(prefix="sdmx_convert_"))
    result = subprocess.run(
        [executable, "--headless", "--convert-to", target_extension.lstrip("."), "--outdir", str(temp_dir), str(path)],
        capture_output=True, text=True, timeout=180,
    )
    converted = temp_dir / f"{path.stem}{target_extension}"
    if result.returncode or not converted.exists():
        raise RuntimeError(f"LibreOffice conversion failed: {result.stderr.strip() or result.stdout.strip()}")
    return converted


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        parts = [f"[DOCUMENT TITLE]\n{path.stem}"]
        with fitz.open(path) as document:
            for number, page in enumerate(document, 1):
                parts.append(f"[PAGE {number}]\n{page.get_text('text', sort=True).strip()}")
        return "\n\n".join(parts)
    if suffix in {".doc", ".ppt"}:
        converted = _office_legacy(path, ".docx" if suffix == ".doc" else ".pptx")
        try:
            return extract_text(converted)
        finally:
            shutil.rmtree(converted.parent, ignore_errors=True)
    if suffix == ".docx":
        if not zipfile.is_zipfile(path):
            raise ValueError("Invalid DOCX container")
        document = Document(path)
        parts = [f"[DOCUMENT TITLE]\n{path.stem}"]
        table_number = 0
        for block in document.element.body.iterchildren():
            if block.tag.endswith("}p"):
                text = "".join(node.text or "" for node in block.iter() if node.tag.endswith("}t")).strip()
                if text:
                    style = block.find(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pStyle")
                    marker = "[SECTION]" if style is not None and "Heading" in str(style.attrib) else "[PARAGRAPH]"
                    parts.append(f"{marker}\n{text}")
            elif block.tag.endswith("}tbl"):
                table_number += 1
                rows = [["".join(node.text or "" for node in cell.iter() if node.tag.endswith("}t")) for cell in row if cell.tag.endswith("}tc")] for row in block if row.tag.endswith("}tr")]
                parts.append(f"[TABLE {table_number}]\n{_table(rows)}")
        return "\n\n".join(parts)
    if suffix in {".xlsx", ".xls"}:
        parts = [f"[DOCUMENT TITLE]\n{path.stem}"]
        if suffix == ".xlsx":
            workbook = load_workbook(path, read_only=True, data_only=False)
            for sheet in workbook.worksheets:
                parts.append(f"[WORKSHEET]\n{sheet.title}")
                rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
                parts.append(f"[TABLE]\n{_table(rows)}")
            workbook.close()
        else:
            sheets = pd.read_excel(path, sheet_name=None, header=None, dtype=str)
            for name, frame in sheets.items():
                parts.append(f"[WORKSHEET]\n{name}\n\n[TABLE]\n{_table(frame.fillna('').values.tolist())}")
        return "\n\n".join(parts)
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
            rows = list(csv.reader(handle))
        return f"[DOCUMENT TITLE]\n{path.stem}\n\n[TABLE]\n{_table(rows)}"
    if suffix == ".txt":
        return f"[DOCUMENT TITLE]\n{path.stem}\n\n[TEXT]\n{path.read_text(encoding='utf-8-sig', errors='replace')}"
    if suffix in {".html", ".htm"}:
        soup = BeautifulSoup(path.read_text(encoding="utf-8-sig", errors="replace"), "lxml")
        parts = [f"[DOCUMENT TITLE]\n{(soup.title.string.strip() if soup.title and soup.title.string else path.stem)}"]
        for element in soup.find_all(["h1", "h2", "h3", "p", "table"]):
            if element.name == "table":
                rows = [[cell.get_text(" ", strip=True) for cell in row.find_all(["th", "td"])] for row in element.find_all("tr")]
                parts.append(f"[TABLE]\n{_table(rows)}")
            else:
                value = element.get_text(" ", strip=True)
                if value:
                    parts.append(f"[{'SECTION' if element.name.startswith('h') else 'PARAGRAPH'}]\n{value}")
        return "\n\n".join(parts)
    if suffix == ".pptx":
        if not zipfile.is_zipfile(path):
            raise ValueError("Invalid PPTX container")
        presentation = Presentation(path)
        parts = [f"[DOCUMENT TITLE]\n{path.stem}"]
        for slide_number, slide in enumerate(presentation.slides, 1):
            parts.append(f"[SLIDE {slide_number}]")
            for shape in slide.shapes:
                if getattr(shape, "has_table", False):
                    rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                    parts.append(f"[TABLE]\n{_table(rows)}")
                elif hasattr(shape, "text") and shape.text.strip():
                    parts.append(shape.text.strip())
        return "\n\n".join(parts)
    raise ValueError(f"Unsupported extension: {suffix}")


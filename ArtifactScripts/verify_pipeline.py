#!/usr/bin/env python3
from __future__ import annotations

import json
import shutil
import sys
import tempfile
from dataclasses import replace
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "Codes"))

from pipeline.artifacts import discover_artifacts, file_hash  # noqa: E402
from pipeline.config import ARTIFACT_ROOT, CONTROL_FILE, EMBEDDING_ROOT, INPUTS, TEXT_ROOT  # noqa: E402
import pipeline.control as control_module  # noqa: E402
import pipeline.agents as agent_module  # noqa: E402
from pipeline.models import model_for  # noqa: E402
from run_pipeline import folder_name  # noqa: E402


def evidence_quotes(value: Any):
    if isinstance(value, dict):
        if isinstance(value.get("evidence_quote"), str):
            yield value["evidence_quote"]
        for nested in value.values():
            yield from evidence_quotes(nested)
    elif isinstance(value, list):
        for nested in value:
            yield from evidence_quotes(nested)


def verify_control_invalidation(artifacts, source: Path) -> dict[str, bool]:
    original = control_module.CONTROL_FILE
    temporary_paths: list[Path] = []
    try:
        def control_copy() -> Path:
            handle = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
            handle.close()
            path = Path(handle.name)
            shutil.copy2(CONTROL_FILE, path)
            temporary_paths.append(path)
            return path

        control_module.CONTROL_FILE = control_copy()
        changed = control_module.ControlWorkbook(artifacts)
        changed_row, _ = changed.sync_file(source, source.relative_to(INPUTS).as_posix(), "0" * 64)
        prefixes = ["TXT", "EMBEDDING", *[item.key for item in artifacts]]
        source_change_ok = all(changed.get(changed_row, f"{prefix}_STATUS") == "PENDING" for prefix in prefixes)

        mutated = [replace(artifacts[0], sha256="f" * 64), *artifacts[1:]]
        control_module.CONTROL_FILE = control_copy()
        selective = control_module.ControlWorkbook(mutated)
        selective_row, _ = selective.sync_file(source, source.relative_to(INPUTS).as_posix(), file_hash(source))
        principle_change_ok = selective.get(selective_row, f"{artifacts[0].key}_STATUS") == "PENDING" and all(
            selective.get(selective_row, f"{item.key}_STATUS") == "SUCCESS" for item in artifacts[1:]
        )

        selective.stage(selective_row, artifacts[1].key, "FAILED", "simulated retriable failure")
        failed_retry_ok = selective.get(selective_row, f"{artifacts[1].key}_STATUS") == "FAILED"
        return {"source_change_invalidates_downstream": source_change_ok, "single_principles_change_is_selective": principle_change_ok, "failed_stage_remains_retryable": failed_retry_ok}
    finally:
        control_module.CONTROL_FILE = original
        for path in temporary_paths:
            path.unlink(missing_ok=True)


def verify_schema_repair_loop(artifact, valid_value: dict) -> bool:
    invalid = dict(valid_value)
    invalid.pop("source_report", None)
    responses = [invalid, valid_value]
    prompts: list[str] = []
    original = agent_module.call_structured

    def fake_call(client, settings, prompt, model, schema_name):
        prompts.append(prompt)
        return responses.pop(0)

    try:
        agent_module.call_structured = fake_call
        result = agent_module.generate_and_validate(
            None, {}, artifact, model_for(artifact.key), valid_value["source_report"],
            valid_value["source_sha256"], "synthetic context", [1], "",
        )
        return result.source_report == valid_value["source_report"] and len(prompts) == 2 and "source_report" in prompts[1] and "exact errors" in prompts[1]
    finally:
        agent_module.call_structured = original
def main() -> None:
    source_name = "sample_statistical_report.pdf"
    stem = Path(source_name).stem
    text = (TEXT_ROOT / f"{stem}.txt").read_text(encoding="utf-8")
    required_text = ("[PAGE 1]", "[PAGE 2]", "Table 1. Employment rate", "Table 2. Government expenditure", "Methodology", "Employment rate", "Government expenditure", "Provisional", "Estimated")
    missing = [item for item in required_text if item not in text]
    if missing:
        raise AssertionError(f"TXT structure/content missing: {missing}")

    vectors = np.load(EMBEDDING_ROOT / stem / "embeddings.npy")
    manifest = json.loads((EMBEDDING_ROOT / stem / "manifest.json").read_text(encoding="utf-8"))
    chunks = (EMBEDDING_ROOT / stem / "chunks.jsonl").read_text(encoding="utf-8").splitlines()
    if vectors.shape != (manifest["number_of_chunks"], manifest["embedding_dimension"]) or len(chunks) != vectors.shape[0]:
        raise AssertionError("Embedding manifest/chunks/vector shape mismatch")

    artifacts, _ = discover_artifacts()
    quote_count = 0
    first_artifact_value = None
    for artifact in artifacts:
        path = ARTIFACT_ROOT / folder_name(artifact.key) / f"{stem}.json"
        raw = json.loads(path.read_text(encoding="utf-8"))
        validated = model_for(artifact.key).model_validate(raw)
        if first_artifact_value is None:
            first_artifact_value = validated.model_dump(mode="json")
        quotes = list(evidence_quotes(validated.model_dump(mode="json")))
        unmatched = [quote for quote in quotes if quote not in text]
        if unmatched:
            raise AssertionError(f"{artifact.key} has {len(unmatched)} report evidence quotes absent from TXT")
        quote_count += len(quotes)

    workbook = load_workbook(CONTROL_FILE, read_only=True, data_only=True)
    sheet = workbook["Pipeline Control"]
    headers = [cell.value for cell in sheet[1]]
    rows = [dict(zip(headers, values)) for values in sheet.iter_rows(min_row=2, values_only=True)]
    row = next(item for item in rows if item["FILE_NAME"] == source_name)
    status_columns = ["TXT_STATUS", "EMBEDDING_STATUS", *[f"{artifact.key}_STATUS" for artifact in artifacts]]
    failures = {column: row[column] for column in status_columns if row[column] != "SUCCESS"}
    if failures:
        raise AssertionError(f"Non-success control statuses: {failures}")
    workbook.close()
    invalidation_checks = verify_control_invalidation(artifacts, INPUTS / source_name)
    if not all(invalidation_checks.values()):
        raise AssertionError(f"Control invalidation simulation failed: {invalidation_checks}")
    schema_repair_ok = verify_schema_repair_loop(artifacts[0], first_artifact_value)
    if not schema_repair_ok:
        raise AssertionError("Pydantic schema-repair prompt simulation failed")
    print(json.dumps({
        "source": source_name,
        "txt_characters": len(text),
        "embedding_shape": list(vectors.shape),
        "artifacts_validated": len(artifacts),
        "report_evidence_quotes_verified": quote_count,
        "control_statuses": "all SUCCESS",
        "invalidation_checks": invalidation_checks,
        "schema_repair_loop": schema_repair_ok,
    }, indent=2))


if __name__ == "__main__":
    main()
